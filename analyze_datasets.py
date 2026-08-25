import openpyxl
import pandas as pd

# Deal funnel Data
print("=" * 60)
print("=== Deal funnel Data.xlsx ===")
print("=" * 60)
try:
    wb = openpyxl.load_workbook("C:\\Users\\R Mohith\\Downloads\\Deal funnel Data.xlsx", data_only=True)
    print(f"Sheet names: {wb.sheetnames}")
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\nSheet: {sheet_name}")
        print(f"  Rows: {ws.max_row}, Columns: {ws.max_column}")
        # Print headers
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        print(f"  Headers: {headers}")
        # Print first 5 rows
        for r in range(1, min(ws.max_row + 1, 6)):
            row_data = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            print(f"  Row {r}: {row_data}")
except Exception as e:
    print(f"Error reading Deal funnel: {e}")

# Work Order Tracker
print("\n\n" + "=" * 60)
print("=== Work_Order_Tracker Data.xlsx ===")
print("=" * 60)
try:
    wb2 = openpyxl.load_workbook("C:\\Users\\R Mohith\\Downloads\\Work_Order_Tracker Data.xlsx", data_only=True)
    print(f"Sheet names: {wb2.sheetnames}")
    for sheet_name in wb2.sheetnames:
        ws = wb2[sheet_name]
        print(f"\nSheet: {sheet_name}")
        print(f"  Rows: {ws.max_row}, Columns: {ws.max_column}")
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        print(f"  Headers: {headers}")
        for r in range(1, min(ws.max_row + 1, 6)):
            row_data = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            print(f"  Row {r}: {row_data}")
except Exception as e:
    print(f"Error reading Work Order: {e}")

# Also try pandas
print("\n\n" + "=" * 60)
print("=== Pandas Profiling ===")
print("=" * 60)
try:
    df1 = pd.read_excel("C:\\Users\\R Mohith\\Downloads\\Deal funnel Data.xlsx")
    print(f"\nDeal funnel Data - pandas shape: {df1.shape}")
    print(f"Columns: {list(df1.columns)}")
    print(f"Dtypes:\n{df1.dtypes}")
    print(f"\nFirst 3 rows:\n{df1.head(3).to_string()}")
    print(f"\nMissing values per column:\n{df1.isnull().sum()[df1.isnull().sum() > 0]}")
    print(f"\nDistinct statuses: {df1['Status'].unique() if 'Status' in df1.columns else 'N/A'}")
    print(f"Distinct stages: {df1['Stage'].unique() if 'Stage' in df1.columns else 'N/A'}")
    print(f"Distinct sectors: {df1['Sector'].unique() if 'Sector' in df1.columns else 'N/A'}")
    if 'Close Date' in df1.columns or 'close_date' in df1.columns or 'Closure Date' in df1.columns:
        date_cols = [c for c in df1.columns if 'date' in str(c).lower()]
        print(f"Potential date columns: {date_cols}")
    if 'Probability' in df1.columns or 'probability' in df1.columns or 'closure_prob' in df1.columns:
        prob_cols = [c for c in df1.columns if 'prob' in str(c).lower()]
        print(f"Potential probability columns: {prob_cols}")
    if 'Deal Value' in df1.columns or 'deal_value' in df1.columns or 'Value' in df1.columns:
        val_cols = [c for c in df1.columns if 'value' in str(c).lower() or 'amount' in str(c).lower()]
        print(f"Potential value columns: {val_cols}")

    df2 = pd.read_excel("C:\\Users\\R Mohith\\Downloads\\Work_Order_Tracker Data.xlsx")
    print(f"\nWork Order Tracker - pandas shape: {df2.shape}")
    print(f"Columns: {list(df2.columns)}")
    print(f"Dtypes:\n{df2.dtypes}")
    print(f"\nFirst 3 rows:\n{df2.head(3).to_string()}")
    print(f"\nMissing values per column:\n{df2.isnull().sum()[df2.isnull().sum() > 0]}")
    status_cols = [c for c in df2.columns if any(s in str(c) for s in ['Status', 'status', 'STAGE', 'stage'])]
    print(f"Potential status columns: {status_cols}")
    wo_cols = [c for c in df2.columns if any(s in str(c) for s in ['WO', 'wo', 'Work Order', 'work order'])]
    print(f"Potential WO columns: {wo_cols}")
    billing_cols = [c for c in df2.columns if any(s in str(c) for s in ['Billing', 'billing', 'Bill'])]
    print(f"Potential billing columns: {billing_cols}")
    collection_cols = [c for c in df2.columns if any(s in str(c) for s in ['Collection', 'collection'])]
    print(f"Potential collection columns: {collection_cols}")
    sector_cols = [c for c in df2.columns if 'sector' in str(c).lower()]
    print(f"Potential sector columns: {sector_cols}")
except Exception as e:
    print(f"Error with pandas: {e}")